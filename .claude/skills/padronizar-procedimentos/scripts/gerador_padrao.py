"""
Helpers reutilizaveis para gerar documentos .docx no padrao corporativo
CORP-GQ-ANX-001 (ver references/ na raiz desta skill para as regras que
cada funcao aqui implementa).

Nao e' um script de linha de comando, nem especifico de nenhuma area ou
tipo de documento (ETE, fabricacao, qualidade, etc.) -- e' uma biblioteca
generica para importar e chamar a partir do script de padronizacao de
QUALQUER procedimento que chegar (POP, IT, Manual, Programa, PAC, Tabela
de Higienizacao, PAP...), porque o CONTEUDO (secoes, tabelas, textos)
muda a cada documento -- so a FORMA (fonte, margens, cabecalho, rodape,
capa, campos de pagina) e' fixa para todos e vale a pena reaproveitar.

Uso tipico (os valores abaixo sao so' exemplo -- vem do documento sendo
padronizado a cada execucao, nunca fixos):

    import sys
    sys.path.insert(0, "<caminho-desta-skill>/scripts")
    from gerador_padrao import *

    doc = Document()
    setup_pagina(doc.sections[0])
    build_header(doc.sections[0].header, codigo="<CODIGO-DO-DOCUMENTO>",
                 tipo_documento="<POP|IT|MANUAL|PROGRAMA|...>", titulo="...",
                 revisao="<N>", data_revisao="<DD/MM/AAAA>",
                 data_aprovacao="<DD/MM/AAAA>", unidade="<UNIDADE>",
                 logo_path=LOGO_PADRAO)
    build_footer(doc.sections[0].footer, elaboracao="Fulano", ...)
    build_capa(doc, titulo="...", logo_path=LOGO_PADRAO)  # so' p/ tipos com capa
    nova_secao_apos_capa(doc)  # cabecalho continua, rodape para
    heading(doc, "1. OBJETIVOS")
    para(doc, "texto do objetivo...")
    ...
    salvar_docx(doc, "saida.docx")

Gotchas ja' descobertos na pratica (nao repetir):
- Cada <w:fldChar>/<w:instrText> de um campo (PAGE/NUMPAGES) precisa
  estar em um <w:r> proprio -- ver add_field().
- Nao usar section.different_first_page_header_footer (w:titlePg): leitores
  simples (Quick Look do iOS, por ex.) nao renderizam a variante de
  "primeira pagina". Usar nova_secao_apos_capa() em vez disso.
- python-docx's Image.width/height explodem com ZeroDivisionError se a
  imagem nao tiver DPI informado no arquivo -- normalizar com Pillow antes
  (dpi=(96,96)) se a imagem vier de fonte desconhecida (fotos de celular,
  etc.).
"""
import os
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Arial"
AZUL_DESTAQUE = "DBE5F1"  # cor de destaque padrao da 1a linha de qualquer tabela

LOGO_PADRAO = os.path.join(os.path.dirname(__file__), "..", "assets", "tirolez_logo.jpg")


# ---------------------------------------------------------------------
# Formatacao basica (ver references/formatacao.md)
# ---------------------------------------------------------------------

def set_font(run, size=12, bold=None, italic=None):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rFonts = rpr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rpr.append(rFonts)
    rFonts.set(qn('w:ascii'), FONT)
    rFonts.set(qn('w:hAnsi'), FONT)
    rFonts.set(qn('w:cs'), FONT)
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic


def setup_pagina(section, margem_cm=1.5, distancia_cab_rod_cm=1.25):
    """A4 retrato, margens 1,5cm, distancia cabecalho/rodape 1,25cm."""
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(margem_cm)
    section.bottom_margin = Cm(margem_cm)
    section.left_margin = Cm(margem_cm)
    section.right_margin = Cm(margem_cm)
    section.header_distance = Cm(distancia_cab_rod_cm)
    section.footer_distance = Cm(distancia_cab_rod_cm)


def para(doc_or_cell, text="", size=12, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
         space_after=0, line_spacing=1.5, style=None, first_line_indent=True):
    """Paragrafo de texto corrido: Arial, recuo de 1a linha 1,25cm (regra
    real confirmada no XML do template -- nao vem do estilo Normal)."""
    p = doc_or_cell.add_paragraph(style=style)
    p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing = line_spacing
    if first_line_indent:
        pf.first_line_indent = Cm(1.25)
    if text:
        r = p.add_run(text)
        set_font(r, size=size, bold=bold)
    return p


def heading(doc, numero_e_texto, level=1):
    """Titulo/subtitulo: 12pt, negrito, MAIUSCULO (regra fixa para os 3
    niveis, mesmo quando o estilo Heading do Word nao forcaria negrito)."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.keep_with_next = True
    r = p.add_run(numero_e_texto.upper())
    set_font(r, size=12, bold=True)
    p.paragraph_format.outline_level = level - 1
    return p


def bullet(doc, text, size=12):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run(text)
    set_font(r, size=size)
    return p


# ---------------------------------------------------------------------
# Tabelas (ver references/formatacao.md: 9pt, centralizado, linha 1,15,
# cabecalho de coluna sempre AZUL_DESTAQUE)
# ---------------------------------------------------------------------

def shade_cell(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)


def set_cell_valign_center(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    va = OxmlElement('w:vAlign')
    va.set(qn('w:val'), 'center')
    tcPr.append(va)


def cell_text(cell, text, size=9, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, clear=True):
    if clear:
        cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    lines = text.split("\n")
    r = p.add_run(lines[0])
    set_font(r, size=size, bold=bold)
    for extra in lines[1:]:
        rr = p.add_run()
        rr.add_break()
        r2 = p.add_run(extra)
        set_font(r2, size=size, bold=bold)
    return p


def add_table(doc, rows, header_rows=1, col_widths=None, first_col_bold=False):
    """rows: lista de listas de strings. Primeira(s) linha(s) = cabecalho,
    ganham AZUL_DESTAQUE automaticamente."""
    ncols = len(rows[0])
    t = doc.add_table(rows=len(rows), cols=ncols)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = 'Table Grid'
    if col_widths:
        t.autofit = False
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.cell(ri, ci)
            is_header = ri < header_rows
            bold = is_header or (first_col_bold and ci == 0)
            cell_text(cell, val, size=9, bold=bold)
            if is_header:
                shade_cell(cell, AZUL_DESTAQUE)
            set_cell_valign_center(cell)
            if col_widths:
                cell.width = Cm(col_widths[ci])
    if col_widths:
        for ci, w in enumerate(col_widths):
            for row in t.rows:
                row.cells[ci].width = Cm(w)
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_after = Pt(0)
    return t


def add_photo(doc, image_path, caption, width_cm=9):
    """Figura centralizada + legenda em itálico logo abaixo, tambem
    centralizada (regra: 'Alinhamento de figuras: Centralizado')."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(image_path, width=Cm(width_cm))
    pc = doc.add_paragraph()
    pc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pc.paragraph_format.space_after = Pt(8)
    rc = pc.add_run(caption)
    set_font(rc, size=9, bold=True, italic=True)


# ---------------------------------------------------------------------
# Campo dinamico PAGE / NUMPAGES
# ---------------------------------------------------------------------

def add_field(paragraph, field_code, size=8, cached_value="1"):
    """Cada fldChar/instrText num <w:r> PROPRIO -- empacotar tudo num
    unico run e' schema-valido (XSD nao reclama) mas quebra o parser de
    campo complexo do Word na pratica: o cabecalho inteiro deixa de
    aparecer. Confirmado contra documentos reais da empresa, que sempre
    usam runs separados aqui."""
    r_begin = paragraph.add_run()
    fld_begin = OxmlElement('w:fldChar')
    fld_begin.set(qn('w:fldCharType'), 'begin')
    r_begin._element.append(fld_begin)
    set_font(r_begin, size=size)

    r_instr = paragraph.add_run()
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = field_code
    r_instr._element.append(instr)
    set_font(r_instr, size=size)

    r_sep = paragraph.add_run()
    fld_sep = OxmlElement('w:fldChar')
    fld_sep.set(qn('w:fldCharType'), 'separate')
    r_sep._element.append(fld_sep)
    set_font(r_sep, size=size)

    r_cached = paragraph.add_run(cached_value)
    set_font(r_cached, size=size)

    r_end = paragraph.add_run()
    fld_end = OxmlElement('w:fldChar')
    fld_end.set(qn('w:fldCharType'), 'end')
    r_end._element.append(fld_end)
    set_font(r_end, size=size)


# ---------------------------------------------------------------------
# Cabecalho / rodape / capa (ver references/cabecalho-rodape-capa.md)
# ---------------------------------------------------------------------

def build_header(header, codigo, tipo_documento, titulo, revisao,
                  data_revisao, data_aprovacao, unidade, logo_path=LOGO_PADRAO):
    """Tabela de 1 linha x 3 colunas, cada coluna com paragrafos
    empilhados na mesma celula (nao 6 linhas com merge -- as duas formas
    sao validas mas esta e' mais simples de implementar). Tudo
    centralizado, conforme confirmado no XML do template real."""
    tbl = header.add_table(rows=1, cols=3, width=Cm(17.5))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    widths = [3.5, 10.0, 4.0]  # ~20% / 57% / 23%, igual ao template
    for ci, w in enumerate(widths):
        tbl.rows[0].cells[ci].width = Cm(w)

    logo_cell = tbl.cell(0, 0)
    logo_cell.text = ""
    lp = logo_cell.paragraphs[0]
    lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lr = lp.add_run()
    lr.add_picture(logo_path, width=Cm(2.6))
    set_cell_valign_center(logo_cell)

    title_cell = tbl.cell(0, 1)
    title_cell.text = ""
    title_lines = [(None, None), (tipo_documento, 9), (None, None), (titulo, 14), (None, None)]
    tp0 = title_cell.paragraphs[0]
    tp0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tp0.paragraph_format.space_after = Pt(0)
    for text, size in title_lines[1:]:
        p = title_cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        if text is not None:
            r = p.add_run(text)
            set_font(r, size=size, bold=True)
    set_cell_valign_center(title_cell)

    meta_cell = tbl.cell(0, 2)
    meta_cell.text = ""
    meta = [
        ("Código: ", codigo, False),
        ("Revisão: ", revisao, False),
        ("Data Revisão: ", data_revisao, False),
        ("Data Aprovação: ", data_aprovacao, False),
        ("Página: ", None, False),  # campo dinamico
        ("", unidade, True),
    ]
    mp0 = meta_cell.paragraphs[0]
    for ri, (label, value, value_bold) in enumerate(meta):
        p = mp0 if ri == 0 else meta_cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        if label:
            rl = p.add_run(label)
            set_font(rl, size=8, bold=True)
        if value is not None:
            rv = p.add_run(value)
            set_font(rv, size=8, bold=value_bold)
        else:
            add_field(p, 'PAGE  \\* Arabic  \\* MERGEFORMAT', size=8, cached_value="1")
            rv = p.add_run(' de ')
            set_font(rv, size=8)
            add_field(p, 'NUMPAGES  \\* Arabic  \\* MERGEFORMAT', size=8, cached_value="20")
    set_cell_valign_center(meta_cell)
    return tbl


def build_footer(footer, elaboracao, verificacao, aprovacao,
                  area_elaboracao="", area_verificacao="", area_aprovacao=""):
    """Tabela 3 colunas x 2 linhas (nome + area/cargo), Arial 8pt,
    centralizado -- confirmado em 3 documentos reais com rodape
    populado."""
    tbl = footer.add_table(rows=2, cols=3, width=Cm(17.5))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    for ci in range(3):
        tbl.rows[0].cells[ci].width = Cm(5.7)
        tbl.rows[1].cells[ci].width = Cm(5.7)
    vals = [
        ("Elaboração:", elaboracao, area_elaboracao),
        ("Verificação:", verificacao, area_verificacao),
        ("Aprovação:", aprovacao, area_aprovacao),
    ]
    for ci, (label, name, area) in enumerate(vals):
        cell = tbl.cell(0, ci)
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r1 = p.add_run(label + " ")
        set_font(r1, size=8, bold=True)
        r2 = p.add_run(name)
        set_font(r2, size=8)
        set_cell_valign_center(cell)

        cell2 = tbl.cell(1, ci)
        cell2.text = ""
        p2 = cell2.paragraphs[0]
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(0)
        r3 = p2.add_run(area)
        set_font(r3, size=8)
        set_cell_valign_center(cell2)
    return tbl


def build_capa(doc, titulo, logo_path=LOGO_PADRAO, incluir_imagem=True):
    """Receita real confirmada em 3 documentos (CORP-GQ-PAC-003,
    TIR-GQ-MAN-002, CORP-GQ-PRO-003): parágrafos em branco + titulo
    SOZINHO 28pt negrito centralizado (sem repetir tipo/unidade/codigo
    na capa -- isso ja esta no cabecalho) + imagem centralizada
    (mascote/icone do departamento; usar a logo como alternativa quando
    nao houver arte especifica) + mais paragrafos em branco. So' e'
    obrigatoria para Manuais, PAC, Programas e Procedimentos -- POP/IT
    normalmente NAO tem capa (nao chamar esta funcao para esses tipos)."""
    for _ in range(6):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(titulo)
    set_font(r, size=28, bold=True)

    if incluir_imagem:
        for _ in range(2):
            doc.add_paragraph()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run()
        r.add_picture(logo_path, width=Cm(4))

    for _ in range(8):
        doc.add_paragraph()


def salvar_docx(doc, caminho):
    """Salva corrigindo de antemao um defeito conhecido do python-docx: o
    <w:zoom> que ele gera em word/settings.xml as vezes fica sem o
    atributo w:percent, o que falha validacao de schema estrita (embora o
    Word abra normalmente mesmo assim -- e' so' para o validate.py da
    skill docx nao acusar um erro que na pratica nao trava nada)."""
    settings_el = doc.settings.element
    zoom = settings_el.find(qn('w:zoom'))
    if zoom is not None and zoom.get(qn('w:percent')) is None:
        zoom.set(qn('w:percent'), '100')
    doc.save(caminho)


def nova_secao_apos_capa(doc):
    """Fecha a secao da capa e abre a secao do corpo do documento, COM O
    CABECALHO REPETINDO e o RODAPE vazio dai em diante.

    Nao usar section.different_first_page_header_footer (w:titlePg) para
    isso: e' valido em OOXML e o Word desktop entende, mas leitores mais
    simples (confirmado com o Quick Look do iOS) nao renderizam a
    variante de cabecalho/rodape de "primeira pagina" -- o cabecalho
    inteiro some da capa. Duas secoes reais, cada uma so' com
    header/footer "default" (sem nenhuma variante de primeira pagina),
    e' suportado por qualquer leitor.
    """
    next_section = doc.add_section(WD_SECTION.NEW_PAGE)
    next_section.header.is_linked_to_previous = True  # repete o cabecalho da capa
    next_section.footer.is_linked_to_previous = False
    for p in list(next_section.footer.paragraphs):
        p.text = ""
    return next_section
